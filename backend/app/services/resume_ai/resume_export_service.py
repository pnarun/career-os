"""Export optimized resume content — PDF, DOCX, ATS plain text."""

from __future__ import annotations

import io
import re
from typing import Any

from docx import Document

from app.models.resume import ResumeDocument


def export_resume(
    resume: ResumeDocument,
    *,
    format: str = "txt",
    variant_id: str | None = None,
    optimizations: dict[str, Any] | None = None,
) -> dict[str, Any]:
    """Export resume in requested format."""
    content = _build_optimized_content(resume, variant_id=variant_id, optimizations=optimizations)
    fmt = format.lower().strip()

    if fmt == "docx":
        data = _export_docx(content, resume.filename)
        return {
            "format": "docx",
            "filename": _export_filename(resume.filename, "docx"),
            "content_type": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            "data_base64": data,
            "plain_preview": content[:2000],
        }

    if fmt == "pdf":
        data = _export_pdf(content, resume.filename)
        return {
            "format": "pdf",
            "filename": _export_filename(resume.filename, "pdf"),
            "content_type": "application/pdf",
            "data_base64": data,
            "plain_preview": content[:2000],
        }

    return {
        "format": "txt",
        "filename": _export_filename(resume.filename, "txt"),
        "content_type": "text/plain",
        "content": content,
        "plain_preview": content[:2000],
    }


def _build_optimized_content(
    resume: ResumeDocument,
    *,
    variant_id: str | None,
    optimizations: dict[str, Any] | None,
) -> str:
    lines: list[str] = []
    lines.append("=" * 60)
    lines.append("ATS-OPTIMIZED RESUME EXPORT")
    lines.append("=" * 60)
    lines.append("")

    if variant_id:
        lines.append(f"Variant focus: {variant_id.replace('_', ' ').title()}")
        lines.append("")

    if optimizations and optimizations.get("optimization_plan"):
        lines.append("OPTIMIZATION NOTES (apply truthfully):")
        for item in optimizations["optimization_plan"][:8]:
            lines.append(f"  • [{item.get('priority', 'medium').upper()}] {item.get('action', '')}")
        lines.append("")
        lines.append("-" * 60)
        lines.append("")

    lines.append(resume.raw_text.strip())
    lines.append("")
    lines.append("-" * 60)
    lines.append("SKILLS")
    for skill in resume.skills:
        lines.append(f"  • {skill}")

    return "\n".join(lines)


def _export_filename(original: str, ext: str) -> str:
    base = re.sub(r"\.[^.]+$", "", original) or "resume"
    return f"{base}_optimized.{ext}"


def _export_docx(content: str, original_filename: str) -> str:
    import base64

    doc = Document()
    doc.add_heading("Optimized Resume", level=1)
    for paragraph in content.split("\n"):
        if paragraph.strip():
            doc.add_paragraph(paragraph)
    buffer = io.BytesIO()
    doc.save(buffer)
    return base64.b64encode(buffer.getvalue()).decode("ascii")


def _export_pdf(content: str, original_filename: str) -> str:
    """Minimal PDF export using raw PDF structure (no extra dependencies)."""
    import base64

    lines = content.split("\n")[:80]
    y = 750
    stream_parts = ["BT", "/F1 10 Tf"]
    for line in lines:
        safe = line.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")
        if len(safe) > 90:
            safe = safe[:90] + "..."
        stream_parts.append(f"50 {y} Td ({safe}) Tj")
        stream_parts.append("0 -14 Td")
        y -= 14
        if y < 50:
            break
    stream_parts.append("ET")
    stream = "\n".join(stream_parts)
    stream_bytes = stream.encode("latin-1", errors="replace")

    pdf_parts = [
        b"%PDF-1.4\n",
        b"1 0 obj<< /Type /Catalog /Pages 2 0 R >>endobj\n",
        b"2 0 obj<< /Type /Pages /Kids [3 0 R] /Count 1 >>endobj\n",
        b"3 0 obj<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
        b"/Contents 4 0 R /Resources<< /Font<< /F1 5 0 R >> >> >>endobj\n",
        f"4 0 obj<< /Length {len(stream_bytes)} >>stream\n".encode(),
        stream_bytes,
        b"\nendstream endobj\n",
        b"5 0 obj<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>endobj\n",
        b"xref\n0 6\n0000000000 65535 f \n",
    ]
    body = b"".join(pdf_parts)
    xref_offset = len(body)
    trailer = (
        f"trailer<< /Size 6 /Root 1 0 R >>\nstartxref\n{xref_offset}\n%%EOF"
    ).encode()
    pdf_bytes = body + trailer
    return base64.b64encode(pdf_bytes).decode("ascii")
